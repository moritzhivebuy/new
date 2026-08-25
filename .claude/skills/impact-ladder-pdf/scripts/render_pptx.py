#!/usr/bin/env python3
"""Erzeugt aus demselben JSON wie render.js eine bearbeitbare Praesentation.

Aufruf: python3 scripts/render_pptx.py output/<kunde-slug>-impact-ladder.json

Die Datei laesst sich in Google Drive hochladen und dort als Google-Slides-
Praesentation oeffnen und bearbeiten. Das PDF aus template/template.html bleibt
die verbindliche Fassung fuer den Kunden, diese Datei ist die Arbeitsversion.
Schriften: Frank Ruhl Libre und DM Sans, beide in Google Slides verfuegbar.
"""

import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

DARK = RGBColor(0x20, 0x45, 0x40)
TEAL = RGBColor(0x22, 0x8C, 0x7D)
NEON = RGBColor(0xF1, 0xFF, 0x45)
OFFWHITE = RGBColor(0xFB, 0xFB, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE3, 0xE9, 0xE7)
MUTED = RGBColor(0x3E, 0x5B, 0x56)

SERIF = "Frank Ruhl Libre"
SANS = "DM Sans"

CM = 360000
W = round(33.867 * CM)  # 16:9, 33.87 cm x 19.05 cm
H = round(19.05 * CM)
MARGIN = round(1.9 * CM)


def slugify(value):
    value = (value or "").lower()
    for src, dst in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        value = value.replace(src, dst)
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value or "kunde"


def box(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size, color, *, font=SANS, bold=False, space_after=0,
         space_before=0, first=False, line_spacing=1.25):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color
    return p


def rect(slide, left, top, width, height, fill, *, line=None, line_w=1):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left), Emu(top),
                                   Emu(width), Emu(height))
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    shape.text_frame.word_wrap = True
    return shape


def plain_rect(slide, left, top, width, height, fill):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(left), Emu(top),
                                   Emu(width), Emu(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def blank(prs, background=OFFWHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    plain_rect(slide, 0, 0, W, H, background)
    return slide


def section_label(slide, text, top):
    tf = box(slide, MARGIN, top, W - 2 * MARGIN, round(0.8 * CM))
    para(tf, text.upper(), 12, TEAL, bold=True, first=True)
    return top + round(1.1 * CM)


def slide_title(prs, data):
    slide = blank(prs, DARK)
    plain_rect(slide, 0, 0, round(0.5 * CM), H, NEON)
    tf = box(slide, MARGIN, round(4.2 * CM), W - 2 * MARGIN, round(1 * CM))
    para(tf, "IMPACT LADDER", 13, NEON, bold=True, first=True)
    tf = box(slide, MARGIN, round(5.6 * CM), round(24 * CM), round(6 * CM))
    para(tf, "Was Hivebuy für %s bewirkt" % data["customer"]["name"], 40, WHITE,
         font=SERIF, first=True, line_spacing=1.1)
    tf = box(slide, MARGIN, round(12.4 * CM), round(24 * CM), round(2.4 * CM))
    para(tf, "Erstellt für %s" % data["customer"]["ansprechpartner"], 14, WHITE, first=True)
    para(tf, "Basierend auf unserem Gespräch am %s" % data["customer"]["gespraechsdatum"],
         14, WHITE, space_before=3)
    para(tf, "Stand: %s" % data["date"], 12, BORDER, space_before=3)
    return slide


def slide_potenziale(prs, data):
    slide = blank(prs)
    top = section_label(slide, "Aktuelle Optimierungspotenziale", round(1.6 * CM))
    items = data.get("potenziale", [])
    gap = round(0.6 * CM)
    height = round(2.9 * CM)
    for i, item in enumerate(items[:4]):
        y = top + i * (height + gap)
        card = rect(slide, MARGIN, y, W - 2 * MARGIN, height, WHITE, line=BORDER)
        card.text_frame.word_wrap = True
        plain_rect(slide, MARGIN, y, round(0.16 * CM), height, TEAL)
        tf = box(slide, MARGIN + round(0.7 * CM), y + round(0.55 * CM),
                 W - 2 * MARGIN - round(1.4 * CM), height - round(1 * CM))
        para(tf, item["titel"], 16, DARK, bold=True, first=True)
        para(tf, item["beschreibung"], 13, MUTED, space_before=4)
    return slide


def slide_painkiller(prs, data):
    """Wie optimiert Hivebuy den Prozess, Quelle: Abschnitt Pain Killer."""
    items = [t for t in data.get("painkiller", []) if t]
    if not items:
        return None
    slide = blank(prs)
    top = section_label(slide, "Wie optimiert Hivebuy den Prozess", round(1.6 * CM))
    col_w = round((W - 2 * MARGIN - round(0.8 * CM)) / 2)
    card_h = H - top - round(1.6 * CM)
    half = (len(items) + 1) // 2
    for i, column in enumerate((items[:half], items[half:])):
        if not column:
            continue
        x = MARGIN + i * (col_w + round(0.8 * CM))
        rect(slide, x, top, col_w, card_h, WHITE, line=BORDER)
        plain_rect(slide, x, top, round(0.16 * CM), card_h, TEAL)
        tf = box(slide, x + round(0.9 * CM), top + round(0.9 * CM),
                 col_w - round(1.8 * CM), card_h - round(1.8 * CM))
        for j, bullet in enumerate(column):
            para(tf, "•  " + bullet, 13, DARK, first=(j == 0), space_after=9,
                 line_spacing=1.3)
    return slide


def slide_ladder(prs, data, levels, heading):
    slide = blank(prs)
    top = section_label(slide, heading, round(1.6 * CM))
    col_w = round((W - 2 * MARGIN - round(0.8 * CM)) / 2)
    card_h = H - top - round(1.6 * CM)
    for i, level in enumerate(levels):
        x = MARGIN + i * (col_w + round(0.8 * CM))
        rect(slide, x, top, col_w, card_h, WHITE, line=BORDER)
        plain_rect(slide, x, top, round(0.16 * CM), card_h, TEAL)
        badge = rect(slide, x + round(0.7 * CM), top + round(0.7 * CM),
                     round(1.1 * CM), round(1.1 * CM), DARK)
        btf = badge.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        brun = bp.add_run()
        brun.text = str(level["nummer"])
        brun.font.size = Pt(14)
        brun.font.bold = True
        brun.font.name = SANS
        brun.font.color.rgb = NEON
        tf = box(slide, x + round(2.1 * CM), top + round(0.8 * CM),
                 col_w - round(2.8 * CM), round(1.6 * CM))
        para(tf, level["ebene"], 17, DARK, font=SERIF, first=True, line_spacing=1.1)
        tf = box(slide, x + round(0.7 * CM), top + round(2.8 * CM),
                 col_w - round(1.4 * CM), card_h - round(5.2 * CM))
        for j, bullet in enumerate(level.get("bullets", [])):
            para(tf, "•  " + bullet, 13, DARK, first=(j == 0), space_after=7,
                 line_spacing=1.3)
        chip = rect(slide, x + round(0.7 * CM), top + card_h - round(2.2 * CM),
                    round(2.2 * CM), round(0.7 * CM), NEON)
        ctf = chip.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        crun = cp.add_run()
        crun.text = "IMPACT"
        crun.font.size = Pt(9)
        crun.font.bold = True
        crun.font.name = SANS
        crun.font.color.rgb = DARK
        tf = box(slide, x + round(3.2 * CM), top + card_h - round(2.25 * CM),
                 col_w - round(4 * CM), round(1.6 * CM))
        para(tf, level["impact"], 12, MUTED, bold=True, first=True, line_spacing=1.25)
    return slide


def slide_summary(prs, data):
    slide = blank(prs)
    top = section_label(slide, "Zusammenfassung", round(1.6 * CM))
    card = rect(slide, MARGIN, top, W - 2 * MARGIN, round(4.4 * CM), DARK)
    plain_rect(slide, MARGIN, top, round(0.2 * CM), round(4.4 * CM), NEON)
    tf = box(slide, MARGIN + round(1 * CM), top + round(0.8 * CM),
             W - 2 * MARGIN - round(2 * CM), round(3 * CM))
    para(tf, data["executive_summary"], 16, WHITE, font=SERIF, first=True,
         line_spacing=1.3)
    card.text_frame.word_wrap = True

    top = top + round(5.2 * CM)
    col_w = round((W - 2 * MARGIN - round(1.2 * CM)) / 2)
    zahlen = data.get("zahlenbasis", [])
    annahmen = data.get("annahmen", [])
    for i, (heading, items) in enumerate((
        ("Zahlenbasis aus dem Gespräch", ["%s (%s)" % (z["text"], z["status"]) for z in zahlen]),
        ("Annahmen und offene Punkte", list(annahmen)),
    )):
        if not items:
            continue
        x = MARGIN + i * (col_w + round(1.2 * CM))
        tf = box(slide, x, top, col_w, round(0.7 * CM))
        para(tf, heading.upper(), 11, TEAL, bold=True, first=True)
        tf = box(slide, x, top + round(0.9 * CM), col_w, round(4.6 * CM))
        for j, item in enumerate(items):
            para(tf, "•  " + item, 12, MUTED, first=(j == 0), space_after=5,
                 line_spacing=1.25)

    nxt = data.get("naechster_schritt") or {}
    if nxt:
        y = H - round(3.2 * CM)
        rect(slide, MARGIN, y, W - 2 * MARGIN, round(1.9 * CM), WHITE, line=BORDER)
        tf = box(slide, MARGIN + round(0.8 * CM), y + round(0.55 * CM),
                 W - 2 * MARGIN - round(9 * CM), round(1.2 * CM))
        para(tf, nxt.get("text", ""), 13, DARK, bold=True, first=True)
        cta = rect(slide, W - MARGIN - round(7 * CM), y + round(0.45 * CM),
                   round(6.2 * CM), round(1 * CM), NEON)
        ctf = cta.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        run = cp.add_run()
        run.text = nxt.get("cta", "Folgetermin vereinbaren")
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.name = SANS
        run.font.color.rgb = DARK
        run.hyperlink.address = nxt.get("link", "https://www.hivebuy.com/kontakt")
    return slide


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Aufruf: python3 scripts/render_pptx.py <daten.json>")
    src = Path(sys.argv[1])
    data = json.loads(src.read_text(encoding="utf-8"))

    prs = Presentation()
    prs.slide_width = Emu(W)
    prs.slide_height = Emu(H)

    slide_title(prs, data)
    slide_potenziale(prs, data)
    slide_painkiller(prs, data)
    ladder = data.get("ladder", [])
    slide_ladder(prs, data, ladder[:2], "Vorteile für Bedarfsträger und den Einkauf")
    if len(ladder) > 2:
        slide_ladder(prs, data, ladder[2:4],
                     "Vorteile für Finance & Controlling und die Geschäftsführung")
    slide_summary(prs, data)

    out_dir = src.parent
    stamp = data.get("date") or ""
    try:
        day, month, year = stamp.split(".")
        iso = "%s-%s-%s" % (year, month, day)
    except ValueError:
        iso = date.today().isoformat()
    out = out_dir / ("%s-impact-ladder-%s.pptx" % (iso, slugify(data["customer"]["name"])))
    prs.save(out)
    print("PPTX erstellt: %s" % out)


if __name__ == "__main__":
    main()
